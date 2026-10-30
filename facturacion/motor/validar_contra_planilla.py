#!/usr/bin/env python3
"""Contrasta el motor de sesiones contra lo que la planilla factura de verdad.

Uso:
    python3 validar_contra_planilla.py SESIONES.xlsx --anio 2026

La planilla guarda cuantos dias por semana tiene autorizado cada paciente
(Hoja5) pero no cuales. Asi que el motor no puede dar un numero unico: da el
**rango de numeros posibles** para esa cantidad de dias, descontando feriados.
Con eso alcanza para clasificar cada mes facturado:

    explicable   el numero facturado es uno de los posibles
    por debajo   facturo menos que el minimo posible: hay internacion,
                 ausencias o un tope que el motor todavia no tiene cargado
    por encima   facturo mas que el maximo posible: no se puede explicar con
                 el calendario, hay que revisarlo

No modifica la planilla: solo lee.
"""

import argparse
import itertools
import re
import unicodedata
from collections import Counter
from datetime import date

import openpyxl

from feriados import feriados_nacionales, solo_habiles
from modelo import Autorizacion, AutorizacionDia, ObraSocial
from sesiones import generar_sesiones

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
         "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
SIN_TOPE = ObraSocial(id=0, nombre="(sin tope)", tiene_tope=False)


def clave(nombre):
    s = "".join(c for c in unicodedata.normalize("NFD", str(nombre or ""))
                if unicodedata.category(c) != "Mn")
    return " ".join(sorted(re.sub(r"[^A-Z0-9 ]", " ", s.upper()).split()))


def posibles(cantidad_dias, anio, mes, feriados):
    """Todos los totales alcanzables eligiendo `cantidad_dias` de lunes a viernes."""
    totales = set()
    for combo in itertools.combinations(range(1, 6), cantidad_dias):
        aut = Autorizacion(id=0, paciente_id=0, vigencia_desde=date(anio, 1, 1),
                           vigencia_hasta=date(anio, 12, 31),
                           dias=[AutorizacionDia(d) for d in combo])
        totales.add(generar_sesiones(aut, anio, mes, SIN_TOPE,
                                     feriados=feriados).cantidad)
    return totales


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("sesiones")
    ap.add_argument("--anio", type=int, default=2026)
    ap.add_argument("--sin-trasladables", action="store_true",
                    help="excluye los feriados trasladables, que hay que confirmar")
    args = ap.parse_args()

    feriados = solo_habiles(feriados_nacionales(
        args.anio, incluir_trasladables=not args.sin_trasladables))
    print(f"Validando {args.anio} con {len(feriados)} feriados habiles cargados.\n")

    libro = openpyxl.load_workbook(args.sesiones, data_only=True)
    semanales = {clave(r[0]): r[3]
                 for r in libro["Hoja5"].iter_rows(min_row=2, values_only=True)
                 if r[0] and isinstance(r[3], (int, float)) and 1 <= r[3] <= 5}

    cache, resumen = {}, Counter()
    por_debajo, por_encima, sin_autorizacion = [], [], set()

    for nombre_hoja in libro.sheetnames:
        u = nombre_hoja.upper().strip()
        if str(args.anio) not in u:
            continue
        mes = next((i + 1 for i, m in enumerate(MESES) if u.startswith(m)), None)
        if mes is None:
            continue
        ws = libro[nombre_hoja]
        for fila in range(2, ws.max_row + 1):
            paciente = ws.cell(fila, 1).value
            facturadas = ws.cell(fila, 5).value
            if not paciente or not isinstance(facturadas, (int, float)) or not facturadas:
                continue
            k = clave(paciente)
            if k not in semanales:
                sin_autorizacion.add(str(paciente).strip())
                continue
            dias = int(semanales[k])
            if (dias, mes) not in cache:
                cache[(dias, mes)] = posibles(dias, args.anio, mes, feriados)
            rango = cache[(dias, mes)]
            fact = int(facturadas)
            if fact in rango:
                resumen["explicable"] += 1
            elif fact < min(rango):
                resumen["por debajo"] += 1
                por_debajo.append((MESES[mes - 1], str(paciente).strip(), dias,
                                   fact, min(rango)))
            else:
                resumen["por encima"] += 1
                por_encima.append((MESES[mes - 1], str(paciente).strip(), dias,
                                   fact, max(rango)))

    total = sum(resumen.values())
    print(f"{'resultado':14} {'casos':>6}   {'%':>6}")
    for etiqueta in ("explicable", "por debajo", "por encima"):
        n = resumen[etiqueta]
        print(f"{etiqueta:14} {n:>6}   {100 * n / total if total else 0:>5.1f}%")
    print(f"{'TOTAL':14} {total:>6}")

    if por_encima:
        print(f"\nFacturaron mas de lo que permite el calendario ({len(por_encima)}):")
        print(f"  {'mes':12} {'paciente':24} {'d/sem':>5} {'facturo':>7} {'maximo':>7}")
        for m, p, d, f, mx in por_encima[:20]:
            print(f"  {m:12} {p[:24]:24} {d:>5} {f:>7} {mx:>7}")

    if por_debajo:
        print(f"\nFacturaron menos del minimo ({len(por_debajo)}): internacion, "
              f"ausencias o un tope sin cargar")
        print(f"  {'mes':12} {'paciente':24} {'d/sem':>5} {'facturo':>7} {'minimo':>7}")
        for m, p, d, f, mn in por_debajo[:20]:
            print(f"  {m:12} {p[:24]:24} {d:>5} {f:>7} {mn:>7}")

    if sin_autorizacion:
        print(f"\nSin dias semanales cargados en Hoja5 ({len(sin_autorizacion)}): "
              f"no se pueden validar")
        for p in sorted(sin_autorizacion)[:15]:
            print(f"  {p}")


if __name__ == "__main__":
    main()
